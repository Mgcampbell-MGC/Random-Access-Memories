import pandas as pd, json, re, sys, datetime as dt
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import FormulaRule, CellIsRule
from openpyxl.utils import get_column_letter as L
D=pd.read_pickle('final.pkl')
meta=json.load(open('meta.json')) if len(sys.argv)<2 else {}
OUT=sys.argv[1] if len(sys.argv)>1 else 'O_LANCAMENTO_CRM.xlsx'
wb=Workbook()
P=wb.active; P.title='Painel'
S=wb.create_sheet('Lista')
GREEN='E3F1E4'; AMBER='FFF1D6'; GREY='EEEEEE'; HEAD='1F3A2E'; CRM='7A4E1D'
hfont=Font(bold=True,color='FFFFFF'); thin=Side(style='thin',color='DDDDDD')
STATUS=['Não contatado','Tentei – sem resposta','Conversei','Interessada(o)','Proposta enviada','Fechou','Perdido','Não contatar']
COLS=[('Nº',6,'n'),('Nota',6,'grade'),('Marca',22,'marca'),
 ('Status',20,None),('Próximo passo',26,None),('Data do próximo passo',13,None),('Último contato',13,None),('Notas',30,None),
 ('WhatsApp',17,'whatsapp'),('Telefone',17,'telefone'),('E-mail',30,'email'),('Site',30,'site'),('Instagram',22,'instagram'),('Pessoas (sócios no CNPJ)',34,'pessoas'),
 ('Categoria',16,'categoria'),('Formato',16,'formato'),('Produtos novos (90 dias)',10,'n_new'),('Último registro ANVISA',13,'ultimo'),
 ('Cidade',18,'cidade'),('UF',5,'uf'),('Porte',16,'porte'),('Por que esta nota',60,'motivo'),('Alertas',40,'alertas'),
 ('Gasto no último lançamento (R$)',14,None),('Comprou em 30 dias?',11,None),('Pacote de interesse',18,None),
 ('Origem do WhatsApp',30,'whatsapp_origem'),('Tipo e origem do telefone',34,'telefone_tipo'),('Origem do e-mail',26,'email_origem'),
 ('Outros telefones',40,'outros_tel'),('Outros e-mails',40,'outros_email'),('Empresa dona',34,'empresa'),('CNPJ',19,'cnpj'),('Como o dono foi identificado',40,'como_dono'),
 ('Capital social (R$)',14,'capital'),('Abertura',11,'abertura'),('Exemplos de produtos novos',60,'exemplos'),('Primeiro registro ANVISA',13,'primeiro'),
 ('Fabricante / outros titulares na ANVISA',50,'fabricante'),('Plataforma da loja',14,'plataforma'),('Buscar no Google',12,'busca_google'),('Buscar no Instagram',12,'busca_instagram'),('Processos ANVISA (mais recentes)',40,'processos')]
CI={c[0]:i+1 for i,c in enumerate(COLS)}
crmcols={'Status','Próximo passo','Data do próximo passo','Último contato','Notas','Gasto no último lançamento (R$)','Comprou em 30 dias?','Pacote de interesse'}
for j,(h,w,k) in enumerate(COLS,1):
    c=S.cell(1,j,h); c.font=hfont; c.fill=PatternFill('solid',fgColor=CRM if h in crmcols else HEAD)
    c.alignment=Alignment(wrap_text=True,vertical='center'); S.column_dimensions[L(j)].width=w
S.row_dimensions[1].height=42
link_font=Font(color='0563C1',underline='single')
def todate(x):
    if x is None or (isinstance(x,float) and pd.isna(x)) or x=='': return None
    try: return pd.Timestamp(x).to_pydatetime().date()
    except Exception: return None
for i,r in enumerate(D.itertuples(index=False),2):
    r=r._asdict()
    for j,(h,w,k) in enumerate(COLS,1):
        cell=S.cell(i,j)
        if h=='Status': cell.value='Não contatado'; continue
        if k is None: continue
        v=r.get(k)
        if v is None or (isinstance(v,float) and pd.isna(v)): v=''
        if k in ('ultimo','primeiro','abertura'):
            d=todate(v); cell.value=d; cell.number_format='DD/MM/YYYY'; continue
        if k=='capital':
            try: cell.value=float(v) if v!='' else None; cell.number_format='#,##0'
            except: cell.value=None
            continue
        if k=='whatsapp' and v:
            cell.value=v; cell.hyperlink=r['whatsapp_link']; cell.font=link_font; continue
        if k=='telefone' and v:
            cell.value=v; cell.hyperlink='tel:+55'+re.sub(r'\D','',v); cell.font=link_font; continue
        if k=='email' and v:
            cell.value=v; cell.hyperlink='mailto:'+v; cell.font=link_font; continue
        if k in ('site','instagram') and v:
            cell.value=re.sub(r'^https?://(www\.)?','',v).rstrip('/'); cell.hyperlink=v; cell.font=link_font; continue
        if k in ('busca_google','busca_instagram') and v:
            cell.value='buscar'; cell.hyperlink=v; cell.font=link_font; continue
        cell.value=v
n=len(D)+1
S.freeze_panes='D2'
S.column_dimensions.group(L(CI['Origem do WhatsApp']),L(len(COLS)),outline_level=1,hidden=False)
S.sheet_properties.outlinePr.summaryRight=False
S.auto_filter.ref=f'A1:{L(len(COLS))}{n}'
def dv(col,vals,allow_blank=True):
    d=DataValidation(type='list',formula1='"'+','.join(vals)+'"',allow_blank=allow_blank); S.add_data_validation(d); d.add(f'{L(CI[col])}2:{L(CI[col])}{n+500}')
dv('Status',STATUS); dv('Comprou em 30 dias?',['Sim','Não']); dv('Pacote de interesse',['O LANÇAMENTO','ESSENCIAL + FILME','VITRINE','EXTENSÃO','PRÉVIA','Nenhum'])
for col in ('Data do próximo passo','Último contato'):
    for i in range(2,n+1): S.cell(i,CI[col]).number_format='DD/MM/YYYY'
for i in range(2,n+1): S.cell(i,CI['Gasto no último lançamento (R$)']).number_format='#,##0'
g=L(CI['Nota']); rng=f'A2:{L(len(COLS))}{n}'
S.conditional_formatting.add(f'{g}2:{g}{n}',CellIsRule(operator='equal',formula=['"A"'],fill=PatternFill('solid',fgColor='B9DFBF'),font=Font(bold=True)))
S.conditional_formatting.add(f'{g}2:{g}{n}',CellIsRule(operator='equal',formula=['"B"'],fill=PatternFill('solid',fgColor='FFE0A3')))
S.conditional_formatting.add(f'{g}2:{g}{n}',CellIsRule(operator='equal',formula=['"C"'],fill=PatternFill('solid',fgColor='DDDDDD')))
st=L(CI['Status'])
for val,col in [('Interessada(o)','C6EFCE'),('Proposta enviada','9BD3A6'),('Fechou','63BE7B'),('Perdido','F4C7C3'),('Não contatar','E6B8B7'),('Conversei','DDEBF7'),('Tentei – sem resposta','FCE4D6')]:
    S.conditional_formatting.add(f'{st}2:{st}{n}',CellIsRule(operator='equal',formula=[f'"{val}"'],fill=PatternFill('solid',fgColor=col)))
for i in range(2,n+1):
    for j in (CI['Por que esta nota'],CI['Alertas'],CI['Exemplos de produtos novos'],CI['Pessoas (sócios no CNPJ)']):
        S.cell(i,j).alignment=Alignment(wrap_text=False,vertical='top')
# ---------------- Painel
P.sheet_view.showGridLines=False
P.column_dimensions['A'].width=3
for col,w in zip('BCDEFGHIJ',[34,14,14,14,4,30,16,16,40]): P.column_dimensions[col].width=w
def H(r,c,t,size=12):
    x=P.cell(r,c,t); x.font=Font(bold=True,size=size,color=HEAD); return x
def lab(r,c,t,bold=False): x=P.cell(r,c,t); x.font=Font(bold=bold); return x
LS="Lista"
def col(name): return f"{LS}!${L(CI[name])}$2:${L(CI[name])}${n+500}"
P['B2']='O LANÇAMENTO — marcas lançando produtos agora'; P['B2'].font=Font(bold=True,size=18,color=HEAD)
P['B3']=meta.get('subtitle',''); P['B3'].font=Font(italic=True,color='555555')
r=5; H(r,2,'Visão geral'); r+=1
rows=[('Marcas na lista',f'=COUNTA({col("Marca")})'),('Nota A — ligar primeiro',f'=COUNTIF({col("Nota")},"A")'),('Nota B',f'=COUNTIF({col("Nota")},"B")'),('Nota C',f'=COUNTIF({col("Nota")},"C")'),
      ('Com WhatsApp',f'=COUNTIF({col("WhatsApp")},"?*")'),('Com telefone',f'=COUNTIF({col("Telefone")},"?*")'),('Com e-mail',f'=COUNTIF({col("E-mail")},"?*")'),
      ('Com site encontrado',f'=COUNTIF({col("Site")},"?*")'),('Dono identificado (CNPJ)',f'=COUNTIF({col("CNPJ")},"?*")'),('Formato não testado (maquiagem/perfume)',f'=COUNTIF({col("Formato")},"Não testado*")')]
for t,f in rows:
    lab(r,2,t); c=P.cell(r,3,f); c.number_format='#,##0'; c.font=Font(bold=True); r+=1
r+=1; H(r,2,'Funil (atualiza sozinho pela coluna Status)'); r+=1
for s in STATUS:
    lab(r,2,s); c=P.cell(r,3,f'=COUNTIF({col("Status")},"{s}")'); c.font=Font(bold=True); r+=1
r+=1; H(r,2,'Teste das 20 conversas'); r+=1
lab(r,2,'Regras definidas antes das ligações. Preencha "Gasto no último lançamento" e "Comprou em 30 dias?" a cada conversa.'); P.cell(r,2).font=Font(italic=True,color='555555'); r+=1
conv=r
lab(r,2,'Conversas feitas'); P.cell(r,3,'='+'+'.join(f'COUNTIF({col("Status")},"{s}")' for s in ['Conversei','Interessada(o)','Proposta enviada','Fechou','Perdido'])).font=Font(bold=True); r+=1
g25=r; lab(r,2,'Disseram gasto ≥ R$2.500'); P.cell(r,3,f'=COUNTIF({col("Gasto no último lançamento (R$)")},">=2500")').font=Font(bold=True); r+=1
g10=r; lab(r,2,'Disseram gasto ≥ R$1.000'); P.cell(r,3,f'=COUNTIF({col("Gasto no último lançamento (R$)")},">=1000")').font=Font(bold=True); r+=1
glo=r; lab(r,2,'Disseram gasto < R$1.000'); P.cell(r,3,f'=COUNTIF({col("Gasto no último lançamento (R$)")},"<1000")').font=Font(bold=True); r+=1
buy=r; lab(r,2,'Compraram em 30 dias'); P.cell(r,3,f'=COUNTIF({col("Comprou em 30 dias?")},"Sim")').font=Font(bold=True); r+=1
vit=r; lab(r,2,'Compraram VITRINE'); P.cell(r,3,f'=COUNTIFS({col("Comprou em 30 dias?")},"Sim",{col("Pacote de interesse")},"VITRINE")').font=Font(bold=True); r+=1
lab(r,2,'Resultado',True)
P.cell(r,3,f'=IF(C{conv}<20,"Faltam "&(20-C{conv})&" conversas",IF(AND(C{g25}>=6,OR(C{buy}>=2,C{vit}>=4)),"MANTER R$2.990",IF(C{g10}>=8,"MUDAR: ESSENCIAL R$1.490 + FILME R$990",IF(C{glo}>=12,"PARAR","Sem decisão: rever com o Matthew"))))').font=Font(bold=True,color='7A4E1D',size=12)
r+=1
for t in ['MANTER R$2.990: ≥6 de 20 dizem ≥R$2.500 no último lançamento e ≥2 compram em 30 dias (ou ≥4 VITRINE).','MUDAR para ESSENCIAL R$1.490 + FILME R$990: ≥8 de 20 dizem ≥R$1.000.','PARAR: ≥12 de 20 dizem menos de R$1.000.']:
    lab(r,2,t); P.cell(r,2).font=Font(size=9,color='555555'); r+=1
# right column: categories, UF, notes
rr=5; H(rr,7,'Por categoria'); rr+=1
for cat in ['cabelo','rosto/skincare','corpo','sabonete/corpo','desodorante','maquiagem','perfume/fragrância']:
    P.cell(rr,7,cat); P.cell(rr,8,f'=COUNTIF({col("Categoria")},"{cat}")').font=Font(bold=True); P.cell(rr,9,f'=COUNTIFS({col("Categoria")},"{cat}",{col("Nota")},"A")'); rr+=1
P.cell(5,9,'nota A').font=Font(bold=True,color=HEAD)
rr+=1; H(rr,7,'Por estado (UF do dono)'); rr+=1
for uf,cnt in D[D.uf!=''].uf.value_counts().head(12).items():
    P.cell(rr,7,uf); P.cell(rr,8,f'=COUNTIF({col("UF")},"{uf}")').font=Font(bold=True); P.cell(rr,9,f'=COUNTIFS({col("UF")},"{uf}",{col("Nota")},"A")'); rr+=1
# top 25 static
r=max(r,rr)+2; H(r,2,'As 25 primeiras para ligar (cópia fixa; a lista completa está na aba Lista)'); r+=1
hdr=['Marca','Nota','Produtos novos','WhatsApp','Telefone','Cidade/UF','Por que']
for j,h in enumerate(hdr): c=P.cell(r,2+j if j<4 else 2+j+1 if False else 2+j,h); c.font=hfont; c.fill=PatternFill('solid',fgColor=HEAD)
r+=1
for _,x in D[D.grade=='A'].head(25).iterrows():
    vals=[x.marca,x.grade,x.n_new,x.whatsapp,x.telefone,f'{x.cidade}/{x.uf}' if x.uf else '',x.motivo]
    for j,v in enumerate(vals):
        c=P.cell(r,2+j,v)
        if j==3 and v: c.hyperlink=x.whatsapp_link; c.font=link_font
        if j==4 and v: c.hyperlink='tel:+55'+re.sub(r'\D','',v); c.font=link_font
    r+=1
r+=1; H(r,2,'Como usar'); r+=1
NOTES=meta.get('notes',[])
for t in NOTES:
    c=P.cell(r,2,t); c.alignment=Alignment(wrap_text=True,vertical='top'); P.merge_cells(start_row=r,start_column=2,end_row=r,end_column=10)
    P.row_dimensions[r].height=max(15,15*(1+len(t)//150)); r+=1
wb.save(OUT)
print('saved',OUT,'rows',len(D))
