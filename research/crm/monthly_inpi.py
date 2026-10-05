"""Monthly INPI feed: new beauty, supplement and home-fragrance brands from the trademark gazette.

Runs from scratch in an empty folder. It downloads every gazette published since the last run, finds each
applicant's CNPJ and contacts, drops brands already on the CRM, and writes one .xlsx tab ("Novas marcas")
with the same columns as the CRM's Lista, to be added to the sheet as a new tab.

No contact data is kept in git: the CRM's brands are remembered only as SHA-1 fingerprints in
seen_brand_keys.txt, and the last gazette processed is in inpi_state.json.

    python3 monthly_inpi.py --work /tmp/inpi_work            # all gazettes since the last run
    python3 monthly_inpi.py --work /tmp/inpi_work --from 2909 --to 2912
"""
import argparse, collections, concurrent.futures as cf, datetime as dt, hashlib, html, io, json, os, re, sys
import unicodedata, urllib.parse, xml.etree.ElementTree as ET, zipfile
import pandas as pd, requests

HERE = os.path.dirname(os.path.abspath(__file__))
SEEN = os.path.join(HERE, 'seen_brand_keys.txt')
STATE = os.path.join(HERE, 'inpi_state.json')
UA = {'User-Agent': 'Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/126.0 Safari/537.36',
      'Accept-Language': 'pt-BR,pt;q=0.9'}
CA = '/root/.ccr/ca-bundle.crt' if os.path.exists('/root/.ccr/ca-bundle.crt') else True
ANVISA = 'https://dados.anvisa.gov.br/dados/CONSULTAS/PRODUTOS/'
S = requests.Session()
S.mount('https://', requests.adapters.HTTPAdapter(pool_connections=200, pool_maxsize=200))

# ---------------------------------------------------------------- rules shared with assemble.py
DDD = {'SP': [11,12,13,14,15,16,17,18,19], 'RJ': [21,22,24], 'ES': [27,28], 'MG': [31,32,33,34,35,37,38], 'PR': [41,42,43,44,45,46],
       'SC': [47,48,49], 'RS': [51,53,54,55], 'DF': [61], 'GO': [62,64], 'TO': [63], 'MT': [65,66], 'MS': [67], 'AC': [68], 'RO': [69],
       'BA': [71,73,74,75,77], 'SE': [79], 'PE': [81,87], 'AL': [82], 'PB': [83], 'RN': [84], 'CE': [85,88], 'PI': [86,89],
       'PA': [91,93,94], 'AM': [92,97], 'RR': [95], 'AP': [96], 'MA': [98,99]}
VALID_DDD = {d for v in DDD.values() for d in v}
ACC = re.compile(r'contab|contador|contadores|contadora|escritorio|assessoria|consultoria|consultec|auditoria|fiscal|cont\.|ctb|tribut|legaliza|despachante|societario|admempresas|empresas\d*@|\.cnt\.br|\.adv\.br|advocacia|advogad|juridico', re.I)
JUNK = re.compile(r'typedesign|typefoundry|foundry|fonts?@|@fonts|sentry|wixpress|example|exemplo|seuemail|seu-email|seunome|yourname|youremail|nome@|usuario@|user@|teste@|test@|email@email|meuemail|@domain|@dominio|@site\.|@email\.com|@mail\.com$|noreply|no-reply|webmaster@|postmaster@|privacy@shopify|@sentry', re.I)
IGJUNK = set('wix shopify google sephora facebook instagram meta tiktok youtube nuvemshop tiendanube tray vtex lojaintegrada wordpress woocommerce yampi bagy magazord p reel reels explore accounts stories tv direct _u sharer share'.split())
BIG = re.compile(r"L.?OREAL|UNILEVER|PROCTER|NATURA COSM|NATURA IND|\bAVON\b|BOTIC|BEIERSDORF|JOHNSON|KENVUE|\bCOTY\b|\bELCA\b|REVLON|ELIZABETH ARDEN|GALDERMA|PIERRE FABRE|\bNAOS\b|\bCOSMED\b|HYPERA|CIMED|\bEMS S|EUROFARMA|ACHE LABORAT|MANTECORP|HENKEL|\bKAO\b|SHISEIDO|HINODE|MARY KAY|EMBELLEZE|DEVINTEX|NIELY|COLGATE|RECKITT|BAYER|SANOFI|ABBOTT|GLAXO|HALEON|WELEDA|OCCITANE|RAIA DROGASIL|DROGA RAIA|DROGASIL|DPSP|PAGUE MENOS|\bDIMED\b|DROGARIA ARAUJO|CARREFOUR|AMERICANAS|GRANADO|PHEBO|MEGALABS|LIBBS|FARMOQUIMICA|BIOLAB|\bLEGRAND\b|\bGERMED\b|PRATI.?DONADUZZI|\bMEDLEY\b|\bTEUTO\b|LVMH|\bPUIG\b|CHANEL|\bSAVOY\b|BETTECH|LAPRONAT|JEQUITI|SILVIO SANTOS|INTERPARFUMS|AMOREPACIFIC|ESTEE|SEPHORA|MINISO|NIVEA|PIERRE ALEXANDER|ANA HICKMANN COSMETICOS|DAVENE|FARMAX|SALON LINE|BIOMEDIC|THERASKIN|DERMAGE|ADCOS|BEL COL|FLORA PRODUTOS|\bIFF\b|GIVAUDAN|FIRMENICH|SYMRISE|INDITEX|ZARA HOME|\bH&M\b|LOJAS RENNER|RIACHUELO|C&A MODAS|BOTICARIO|\bNATURA (COSMETICOS|INDUSTRIA)|COSMETICOS NATURA LTDA|NATURA &CO|NESTL[EÉ]|\bAMBEV\b|HEINEKEN|DANONE|BIOLAB SANUS|FLORA PRODUTOS", re.I)

def valid_email(e):
    e = e.strip().lower().strip('.')
    if not re.match(r'^[a-z0-9._%+-]+@([a-z0-9-]+\.)+([a-z]{2,6})$', e) or JUNK.search(e): return False
    return not re.search(r'\d+\.\d+', e.split('@')[1])

def placeholder(p):
    d = re.sub(r'\D', '', p)
    if len(d) < 10: return True
    n = d[2:]
    return len(set(n)) <= 2 or bool(re.search(r'12345|23456|34567|45678|98765|87654|00000|11111|99999', n))

def phone_kind(p):
    p = re.sub(r'\D', '', p)
    if p.startswith('55') and len(p) >= 12: p = p[2:]
    if p.startswith(('800', '300')) and len(p) == 10: p = '0' + p
    if p.startswith('0') and not p.startswith(('0800', '0300')) and len(p) in (11, 12): p = p[1:]
    if p.startswith(('0800', '0300', '4004', '4003')): return p, '0800'
    if len(p) < 10 or int(p[:2]) not in VALID_DDD: return p, '?'
    ddd, num = p[:2], p[2:]
    if len(num) == 9 and num[0] == '9': return p, 'celular'
    if len(num) == 8 and num[0] in '6789': return ddd + '9' + num, 'celular (número antigo, 9 acrescentado)'
    if len(num) == 8 and num[0] in '2345': return p, 'fixo'
    return p, '?'

def fmt(p):
    p = re.sub(r'\D', '', p)
    if p.startswith(('0800', '0300', '4004', '4003')): return f'{p[:4]} {p[4:7]} {p[7:]}'
    if len(p) == 11: return f'({p[:2]}) {p[2:7]}-{p[7:]}'
    if len(p) == 10: return f'({p[:2]}) {p[2:6]}-{p[6:]}'
    return p

def cnpj_dv(c12):
    w1 = [5,4,3,2,9,8,7,6,5,4,3,2]
    def d(b, w):
        r = sum(int(x) * y for x, y in zip(b, w)) % 11
        return '0' if r < 2 else str(11 - r)
    a = d(c12, w1)
    return c12 + a + d(c12 + a, [6] + w1)

def cnpj_ok(c):
    return len(c) == 14 and len(set(c)) > 1 and cnpj_dv(c[:12]) == c and not c.startswith('12345678')

def cnpj_fmt(c): return f'{c[:2]}.{c[2:5]}.{c[5:8]}/{c[8:12]}-{c[12:]}' if c else ''

def bkey(s): return re.sub(r'[^a-z0-9]', '', unicodedata.normalize('NFKD', str(s)).encode('ascii', 'ignore').decode().lower())

def fp(s): return hashlib.sha1(s.encode()).hexdigest()

STOP = r'\b(LTDA|LIMITADA|ME|EPP|EIRELI|S ?A|SA|SLU|UNIPESSOAL|MEI|EM RECUPERACAO JUDICIAL)\b'
def name_key(s):
    s = unicodedata.normalize('NFKD', str(s or '')).encode('ascii', 'ignore').decode().upper()
    s = re.sub(r'[^A-Z0-9 ]', ' ', s); s = re.sub(STOP, ' ', s)
    return re.sub(r'\s+', ' ', s).strip()

GENERIC = r'\b(COSMETICOS?|INDUSTRIA|COMERCIO|DE|DO|DA|E|PRODUTOS|BELEZA|DISTRIBUIDORA|IMPORTACAO|EXPORTACAO|SUPLEMENTOS?)\b'
def name_tokens(s):
    s = re.sub(r'^\d{8}\s', '', name_key(s)); s = re.sub(GENERIC, ' ', s)
    return {t for t in s.split() if len(t) >= 3}

def get(url, t=(5, 60), **kw):
    for a in range(4):
        try:
            r = S.get(url, headers=UA, timeout=t, verify=CA, **kw)
            if r.status_code == 200: return r
            if r.status_code == 404: return r
        except Exception:
            pass
    return None

# ---------------------------------------------------------------- 1. gazettes
def latest_rpi():
    r = get('https://revistas.inpi.gov.br/rpi/')
    nums = [int(x) for x in re.findall(r'RM(\d{4})\.zip', r.text)] if r is not None else []
    if not nums: sys.exit('could not read the INPI gazette index')
    return max(nums)

SUPP = re.compile(r'suplement', re.I); HOME = re.compile(r'vela|difusor|aromatiz|home spray', re.I)
def parse_gazette(n, work):
    path = os.path.join(work, f'RM{n}.zip')
    if not os.path.exists(path):
        r = get(f'https://revistas.inpi.gov.br/txt/RM{n}.zip', t=(10, 600))
        if r is None or r.status_code != 200: print(f'RPI {n}: not available, skipped'); return [], None
        open(path, 'wb').write(r.content)
    z = zipfile.ZipFile(path); name = z.namelist()[0]
    rows, date = [], None
    with z.open(name) as fh:
        for ev, el in ET.iterparse(fh, events=('start', 'end')):
            if ev == 'start' and el.tag == 'revista': date = el.attrib.get('data')
            if ev != 'end' or el.tag != 'processo': continue
            if 'IPAS009' in [d.attrib.get('codigo') for d in el.iter('despacho')]:
                tit = [t.attrib for t in el.iter('titular')]
                m = el.find('marca'); nome = ((m.findtext('nome') if m is not None else None) or '').strip()
                cls = {c.attrib.get('codigo'): (c.findtext('especificacao') or '') for c in el.iter('classe-nice')}
                mk = []
                if '03' in cls: mk.append('Cosméticos')
                if '05' in cls and SUPP.search(cls['05']): mk.append('Suplementos')
                if '04' in cls and HOME.search(cls['04']): mk.append('Casa e aroma')
                if mk and nome and tit and tit[0].get('pais') == 'BR':
                    rows.append(dict(rpi=n, processo=el.attrib.get('numero'), deposito=el.attrib.get('data-deposito'), marca=nome,
                                     titular=tit[0].get('nome-razao-social', '').strip(), uf=tit[0].get('uf', ''), classes=' + '.join(mk)))
            el.clear()
    print(f'RPI {n} ({date}): {len(rows)} filings in the three markets')
    return rows, date

# ---------------------------------------------------------------- 2. CNPJ by name (ANVISA lists) and by MEI name
def anvisa_index(work):
    idx = collections.defaultdict(set)
    for f in ['TA_CONSULTA_COSMETICOS.CSV', 'TA_CONSULTA_ALIMENTOS.CSV', 'TA_CONSULTA_SANEANTES.CSV']:
        p = os.path.join(work, f)
        if not os.path.exists(p):
            r = get(ANVISA + f, t=(10, 900))
            if r is None or r.status_code != 200: print('ANVISA list not available:', f); continue
            open(p, 'wb').write(r.content)
        t = pd.read_csv(p, sep=';', encoding='latin-1', dtype=str, usecols=['NO_RAZAO_SOCIAL_EMPRESA', 'NU_CNPJ_EMPRESA']).drop_duplicates()
        for n, c in t.itertuples(index=False):
            k = name_key(n); c = re.sub(r'\D', '', str(c)).zfill(14)
            if k and cnpj_ok(c): idx[k].add(c)
    return idx

def find_cnpj(titular, idx):
    m = re.match(r'^(\d{2})\.?(\d{3})\.?(\d{3})\s', titular)
    if m: return cnpj_dv(''.join(m.groups()) + '0001'), 'MEI: CNPJ calculado do nome'
    s = idx.get(name_key(titular), set())
    if s:
        mat = sorted(c for c in s if c[8:12] == '0001') or sorted(s)
        return mat[0], 'nome igual ao da empresa na ANVISA'
    return '', ''

# ---------------------------------------------------------------- 3. brand website
KW = ['cabelo','pele','shampoo','skincare','hidratante','cosmétic','cosmetic','beleza','capilar','corporal','sérum','serum','perfume','fragr',
      'maquiagem','sabonete','suplement','colágeno','colageno','vitamina','cápsula','capsula','whey','creatina','vela','difusor','aroma']
BAD = re.compile(r'for sale|à venda|slot|casino|gacor|parked|hostinger|index of|godaddy|domain|domínio', re.I)
CNPJF = re.compile(r'(\d{2}\.?\d{3}\.?\d{3}\s?/?\s?\d{4}\s?-?\s?\d{2})')
GEN = {'Cosméticos': ['cosmeticos','beauty','oficial','store','cosmetics','skin','hair'], 'Suplementos': ['suplementos','nutrition','oficial','store','vitaminas'],
       'Casa e aroma': ['aromas','home','velas','oficial','store']}

def fetch(u):
    try: return S.get(u, headers=UA, timeout=(3, 8), verify=CA, allow_redirects=True)
    except Exception: return None

def valid_site(key, r):
    if r is None or r.status_code >= 400: return None
    t = r.text[:800000]; tl = t.lower()
    m = re.search(r'<title[^>]*>(.*?)</title>', t, re.S | re.I); title = html.unescape(m.group(1).strip()) if m else ''
    host = bkey(re.sub(r'^https?://', '', r.url).split('/')[0])
    if BAD.search(title) or re.search(r'this domain|este dom[ií]nio|buy this domain', tl[:20000]): return None
    br = bool(re.search(r'R\$|lang=["\']pt|pt-br|frete|cnpj|carrinho|sacola|comprar', tl))
    if key in bkey(title + ' ' + host) and sum(1 for k in KW if k in tl) >= 2 and br: return title or host
    return None

def site_contacts(t):
    t = t.replace('\\/', '/')
    u = lambda x: list(dict.fromkeys(x))
    cn = [c for c in (re.sub(r'\D', '', x) for x in CNPJF.findall(t)) if cnpj_ok(c)]
    em = [e.lower() for e in re.findall(r'[\w.+-]+@[\w-]+(?:\.[\w-]+)+', t) if valid_email(e) and not re.search(r'\.(png|jpg|jpeg|webp|gif|svg|js|css)$|wixpress|shopify|nuvemshop|tiendanube|lojaintegrada|vtex|yampi|cloudflare', e, re.I)]
    ig = [x for x in re.findall(r'instagram\.com/([A-Za-z0-9_.]+)', t) if x.lower().strip('.') not in IGJUNK and len(x) > 2]
    wa = re.findall(r'(?:wa\.me/\+?|api\.whatsapp\.com/send/?\?phone=\+?|whatsapp\.com/send\?phone=\+?)(\d{10,13})', t)
    tel = [re.sub(r'\D', '', x) for x in re.findall(r'tel:\+?([\d\s\-\(\)]{10,20})', t)]
    return dict(cnpjs=u(cn)[:4], emails=u(em)[:6], ig=u(ig)[:3], wa=u(wa)[:3], tel=u(tel)[:3])

def find_site(marca, classes):
    key = bkey(marca)
    if len(key) < 4: return None
    gens = sum((GEN.get(m.strip(), []) for m in classes.split('+')), [])
    slugs = list(dict.fromkeys([key] + [key + g for g in gens] + ['use' + key, 'loja' + key]))
    urls = [p + s + t + '/' for s in slugs for t in ['.com.br', '.com'] for p in ['https://', 'https://www.']]
    with cf.ThreadPoolExecutor(16) as ix: got = list(ix.map(fetch, urls))
    for r in got:
        ti = valid_site(key, r)
        if ti:
            base = re.match(r'https?://[^/]+', r.url).group(0); allt = r.text
            for p in ['/pages/contato', '/contato', '/fale-conosco', '/quem-somos', '/policies/terms-of-service', '/politica-de-privacidade']:
                if CNPJF.search(allt) and '@' in allt: break
                rr = fetch(base + p)
                if rr is not None and rr.status_code == 200 and len(rr.text) < 3000000: allt += rr.text
            return dict(url=r.url, **site_contacts(allt))
    return None

# ---------------------------------------------------------------- 4. company register (OpenCNPJ)
REG = {}
def reg(c):
    if not c: return None
    if c in REG: return REG[c]
    r = get(f'https://api.opencnpj.org/{c}', t=(5, 30))
    x = r.json() if (r is not None and r.status_code == 200) else None
    if not x: REG[c] = None; return None
    phones = [f"{t.get('ddd','')}{t.get('numero','')}" for t in x.get('telefones') or [] if not t.get('is_fax')]
    try: cap = float(str(x.get('capital_social')).replace('.', '').replace(',', '.'))
    except Exception: cap = None
    REG[c] = dict(razao=x.get('razao_social', ''), porte=x.get('porte_empresa', '') or '', capital=cap, abertura=x.get('data_inicio_atividade', ''),
                  cidade=(x.get('municipio') or '').title(), uf=x.get('uf', ''), situacao=x.get('situacao_cadastral', ''),
                  socios=[(q.get('nome_socio', ''), q.get('qualificacao_socio', '')) for q in x.get('QSA') or []],
                  emails=[x['email'].lower()] if x.get('email') else [], phones=phones)
    return REG[c]

# ---------------------------------------------------------------- 5. one row per brand, same rules as the CRM
def build_row(r, rpi_date, idx):
    mk = 'Beleza' if 'Cosméticos' in r['classes'] else ('Suplementos' if 'Suplementos' in r['classes'] else 'Casa e aroma')
    owner, how = find_cnpj(r['titular'], idx)
    if owner: how = 'pedido de marca no INPI + ' + how
    i = find_site(r['marca'], r['classes'])
    if i:
        for c in i['cnpjs'][:2]:
            rc = reg(c)
            if rc and rc['razao'] and name_tokens(rc['razao']) & name_tokens(r['titular']):
                if not owner: owner, how = c, 'CNPJ no rodapé do site = titular do pedido no INPI'
                break
    O = reg(owner) if owner else None
    alerts = ['marca nova: pedido de registro no INPI, provavelmente ainda não lançou — ofereça a PRÉVIA']
    wa, tel, emails, regmails = [], [], [], []
    if i:
        for w in i['wa']:
            p, k = phone_kind(w)
            if k not in ('?', '0800') and not placeholder(p): wa.append((fmt(p), 'WhatsApp no site da marca', p))
        for t in i['tel']:
            p, k = phone_kind(t)
            if k != '?' and not placeholder(p): tel.append((fmt(p), k + ' (site da marca)', p))
        emails = [(e, 'site da marca') for e in i['emails'] if not ACC.search(e)]
    if O:
        for p0 in O['phones']:
            if placeholder(p0): continue
            q, k = phone_kind(p0)
            if k == '?': continue
            flag = ' — DDD de outro estado (pode ser do contador)' if O['uf'] in DDD and int(q[:2]) not in DDD[O['uf']] else ''
            tel.append((fmt(q), k + ' (cadastro CNPJ)' + flag, q))
        for a in O['emails']:
            if valid_email(a): regmails.append((a, 'cadastro CNPJ' + (' — PROVÁVEL CONTADOR' if ACC.search(a) else ''), bool(ACC.search(a))))
    wa_best = wa[0] if wa else None
    mob = [t for t in tel if t[1].startswith('celular')]
    if not wa_best and mob: wa_best = (mob[0][0], 'celular do cadastro (pode ter WhatsApp)', mob[0][2])
    tel_sorted = sorted(tel, key=lambda t: ('outro estado' in t[1], 'site' not in t[1]))
    tel_best = next((t for t in tel_sorted if not wa_best or t[2] != wa_best[2]), tel_sorted[0] if tel_sorted else None)
    em_all = emails + [(a, t) for a, t, acc in regmails if not acc]; em_best = em_all[0] if em_all else None
    em_acc = [(a, t) for a, t, acc in regmails if acc]
    if em_acc and not em_best: alerts.append('só há e-mail do contador')
    porte = (O or {}).get('porte', '').upper()
    porte = {'MICROEMPRESA (ME)': 'MICRO EMPRESA', 'EMPRESA DE PEQUENO PORTE (EPP)': 'EMPRESA DE PEQUENO PORTE'}.get(porte, porte)
    large = 'DEMAIS' in porte
    if large: alerts.append('empresa de porte maior (DEMAIS)')
    if O and O['situacao'] and O['situacao'].upper() != 'ATIVA': alerts.append('CNPJ ' + O['situacao'])
    contact = bool(wa_best or tel_best or em_best)
    word = {'Beleza': 'cosméticos', 'Suplementos': 'suplementos', 'Casa e aroma': 'velas aromas'}[mk]
    reason = f"marca nova: pedido de registro no INPI em {r['deposito']} (publicado na RPI {r['rpi']}, {rpi_date}); classe: {r['classes']}"
    reason += ('; dono: ' + how) if how else ('' if contact else '; contato não encontrado: use os links de busca')
    marca = r['marca'].title() if r['marca'].isupper() else r['marca']
    return dict(grade='B' if contact and not large else 'C',
                score=6 + (2 if wa else 0) + (1 if i else 0) + (1 if owner else 0) + (1 if (O or {}).get('uf') == 'SP' else 0) - (3 if large else 0),
                marca=marca, mercado=mk, fonte=f"INPI: pedido de marca (RPI {r['rpi']})", data_inpi=dt.datetime.strptime(r['deposito'], '%d/%m/%Y').date(),
                site=(i or {}).get('url', ''), instagram=('https://instagram.com/' + i['ig'][0]) if i and i['ig'] else '',
                whatsapp=wa_best[0] if wa_best else '', whatsapp_link=('https://wa.me/55' + wa_best[2]) if wa_best else '', whatsapp_origem=wa_best[1] if wa_best else '',
                telefone=tel_best[0] if tel_best else '', telefone_tipo=tel_best[1] if tel_best else '', outros_tel='; '.join(f'{t[0]} [{t[1]}]' for t in tel_sorted[1:5]),
                email=em_best[0] if em_best else '', email_origem=em_best[1] if em_best else '', outros_email='; '.join(f'{a} [{t}]' for a, t in em_all[1:4] + em_acc[:2]),
                pessoas='; '.join(f'{n.title()} ({q})' for n, q in (O or {}).get('socios', [])[:4]), empresa=(O or {}).get('razao') or r['titular'],
                cnpj=cnpj_fmt(owner), como_dono=how or 'titular do pedido de marca no INPI (CNPJ não encontrado)', porte=porte,
                capital=(O or {}).get('capital'), abertura=(O or {}).get('abertura', ''), cidade=(O or {}).get('cidade', ''), uf=(O or {}).get('uf') or r['uf'],
                alertas='; '.join(alerts), motivo=reason, processos=f"INPI {r['processo']}",
                busca_google='https://www.google.com/search?q=' + urllib.parse.quote(f'"{r["marca"]}" {word}'),
                busca_instagram='https://www.google.com/search?q=' + urllib.parse.quote(f'site:instagram.com "{r["marca"]}"'),
                _bk=bkey(r['marca']), _cnpj=owner)

# ---------------------------------------------------------------- 6. workbook tab with the CRM's Lista columns
COLS = [('Nº', 6, 'n'), ('Nota', 6, 'grade'), ('Marca', 22, 'marca'), ('Status', 20, None), ('Próximo passo', 26, None), ('Data do próximo passo', 13, None),
        ('Último contato', 13, None), ('Notas', 30, None), ('WhatsApp', 17, 'whatsapp'), ('Telefone', 17, 'telefone'), ('E-mail', 30, 'email'), ('Site', 30, 'site'),
        ('Instagram', 22, 'instagram'), ('Pessoas (sócios no CNPJ)', 34, 'pessoas'), ('Mercado', 13, 'mercado'), ('Categoria', 16, None), ('Formato', 16, None),
        ('Fonte', 24, 'fonte'), ('Produtos novos (90 dias)', 10, None), ('Último registro ANVISA', 13, None), ('Pedido de marca no INPI', 13, 'data_inpi'),
        ('Cidade', 18, 'cidade'), ('UF', 5, 'uf'), ('Porte', 16, 'porte'), ('Por que esta nota', 60, 'motivo'), ('Alertas', 40, 'alertas'),
        ('Gasto no último lançamento (R$)', 14, None), ('Comprou em 30 dias?', 11, None), ('Pacote de interesse', 18, None),
        ('Origem do WhatsApp', 30, 'whatsapp_origem'), ('Tipo e origem do telefone', 34, 'telefone_tipo'), ('Origem do e-mail', 26, 'email_origem'),
        ('Outros telefones', 40, 'outros_tel'), ('Outros e-mails', 40, 'outros_email'), ('Empresa dona', 34, 'empresa'), ('CNPJ', 19, 'cnpj'),
        ('Como o dono foi identificado', 40, 'como_dono'), ('Capital social (R$)', 14, 'capital'), ('Abertura', 11, 'abertura'),
        ('Exemplos de produtos novos', 60, None), ('Primeiro registro ANVISA', 13, None), ('Fabricante / outros titulares na ANVISA', 50, None),
        ('Plataforma da loja', 14, None), ('Buscar no Google', 12, 'busca_google'), ('Buscar no Instagram', 12, 'busca_instagram'),
        ('Processos ANVISA (mais recentes)', 40, 'processos')]
STATUS = ['Não contatado', 'Tentei – sem resposta', 'Conversei', 'Interessada(o)', 'Proposta enviada', 'Fechou', 'Perdido', 'Não contatar']

def write_xlsx(D, out, title):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.worksheet.datavalidation import DataValidation
    from openpyxl.formatting.rule import CellIsRule
    from openpyxl.utils import get_column_letter as L
    wb = Workbook(); ws = wb.active; ws.title = title
    crm = {'Status', 'Próximo passo', 'Data do próximo passo', 'Último contato', 'Notas', 'Gasto no último lançamento (R$)', 'Comprou em 30 dias?', 'Pacote de interesse'}
    link = Font(color='0563C1', underline='single')
    for j, (h, w, k) in enumerate(COLS, 1):
        c = ws.cell(1, j, h); c.font = Font(bold=True, color='FFFFFF'); c.fill = PatternFill('solid', fgColor='7A4E1D' if h in crm else '1F3A2E')
        c.alignment = Alignment(wrap_text=True, vertical='center'); ws.column_dimensions[L(j)].width = w
    ws.row_dimensions[1].height = 42
    for i, r in enumerate(D.to_dict('records'), 2):
        for j, (h, w, k) in enumerate(COLS, 1):
            cell = ws.cell(i, j)
            if h == 'Status': cell.value = 'Não contatado'; continue
            v = r.get(k) if k else None
            if v is None or (isinstance(v, float) and pd.isna(v)) or v == '': continue
            if k in ('data_inpi', 'abertura'):
                try: cell.value = pd.Timestamp(v).date(); cell.number_format = 'DD/MM/YYYY'
                except Exception: pass
                continue
            cell.value = v
            if k == 'capital': cell.number_format = '#,##0'
            if k == 'whatsapp': cell.hyperlink = r['whatsapp_link']; cell.font = link
            if k == 'telefone': cell.hyperlink = 'tel:+55' + re.sub(r'\D', '', v); cell.font = link
            if k == 'email': cell.hyperlink = 'mailto:' + v; cell.font = link
            if k in ('site', 'instagram'): cell.value = re.sub(r'^https?://(www\.)?', '', v).rstrip('/'); cell.hyperlink = v; cell.font = link
            if k in ('busca_google', 'busca_instagram'): cell.value = 'buscar'; cell.hyperlink = v; cell.font = link
    n = len(D) + 1; CI = {c[0]: i + 1 for i, c in enumerate(COLS)}
    ws.freeze_panes = 'D2'; ws.auto_filter.ref = f'A1:{L(len(COLS))}{n}'
    for col, vals in [('Status', STATUS), ('Comprou em 30 dias?', ['Sim', 'Não']), ('Pacote de interesse', ['O LANÇAMENTO', 'ESSENCIAL + FILME', 'VITRINE', 'EXTENSÃO', 'PRÉVIA', 'Nenhum'])]:
        d = DataValidation(type='list', formula1='"' + ','.join(vals) + '"', allow_blank=True); ws.add_data_validation(d)
        d.add(f'{L(CI[col])}2:{L(CI[col])}{n + 500}')
    g = L(CI['Nota'])
    for val, color in [('A', 'B9DFBF'), ('B', 'FFE0A3'), ('C', 'DDDDDD')]:
        ws.conditional_formatting.add(f'{g}2:{g}{n}', CellIsRule(operator='equal', formula=[f'"{val}"'], fill=PatternFill('solid', fgColor=color)))
    wb.save(out)

# ---------------------------------------------------------------- main
def main():
    ap = argparse.ArgumentParser()
    ap.add_argument('--work', required=True); ap.add_argument('--from', dest='first', type=int); ap.add_argument('--to', dest='last', type=int)
    ap.add_argument('--dry-run', action='store_true', help='do not update seen_brand_keys.txt or inpi_state.json')
    a = ap.parse_args(); os.makedirs(a.work, exist_ok=True)
    state = json.load(open(STATE)) if os.path.exists(STATE) else {'last_rpi': None}
    last = a.last or latest_rpi()
    first = a.first or ((state['last_rpi'] or last - 1) + 1)
    if first > last: print(f'nothing new: last processed RPI {state["last_rpi"]}, latest {last}'); return
    seen = set(open(SEEN).read().split()) if os.path.exists(SEEN) else set()
    filings, dates = [], {}
    for n in range(first, last + 1):
        rows, date = parse_gazette(n, a.work); filings += rows; dates[n] = date
    F = pd.DataFrame(filings)
    if F.empty: print('no filings'); return
    F = F[~F.titular.str.contains(BIG)]
    F['_bk'] = F.marca.map(bkey)
    F = F.sort_values('rpi').drop_duplicates(['_bk', 'titular'], keep='last')
    new = F[~F._bk.map(lambda k: fp('b:' + k)).isin(seen) & (F._bk.str.len() >= 2)]
    print(f'{len(F)} brands after filters; {len(F) - len(new)} already on the CRM; {len(new)} new')
    if new.empty:
        if not a.dry_run:
            json.dump({'last_rpi': last, 'last_run': dt.date.today().isoformat(), 'last_file': None}, open(STATE, 'w'), indent=1)
        return
    idx = anvisa_index(a.work); print('ANVISA name index', len(idx))
    rows = []
    with cf.ThreadPoolExecutor(6) as ex:
        for row in ex.map(lambda r: build_row(r, dates.get(r['rpi']), idx), new.to_dict('records')): rows.append(row)
    D = pd.DataFrame(rows)
    D['_known'] = D._cnpj.map(lambda c: bool(c) and fp('c:' + c) in seen)
    D.loc[D._known, 'alertas'] = D.loc[D._known, 'alertas'] + '; a empresa já tem outra marca na sua lista'
    D['_m'] = D.mercado.map({'Beleza': 0, 'Suplementos': 1, 'Casa e aroma': 2}); D['_g'] = D.grade.map({'B': 1, 'C': 2})
    D = D.sort_values(['_m', '_g', 'score'], ascending=[True, True, False]).reset_index(drop=True)
    D.insert(0, 'n', range(1, len(D) + 1))
    month = dt.date.today().strftime('%Y-%m')
    out = os.path.join(a.work, f'O_LANCAMENTO_novas_marcas_{month}_RPI{first}-{last}.xlsx')
    write_xlsx(D, out, f'Novas marcas {month}')
    print('saved', out)
    print(D.groupby(['mercado', 'grade']).size().to_string())
    print('with CNPJ', (D.cnpj != '').sum(), ' site', (D.site != '').sum(), ' whatsapp', (D.whatsapp != '').sum(),
          ' phone', (D.telefone != '').sum(), ' email', (D.email != '').sum(), ' company already on the CRM', int(D._known.sum()))
    if not a.dry_run:
        seen |= {fp('b:' + k) for k in D._bk} | {fp('c:' + c) for c in D._cnpj if c}
        open(SEEN, 'w').write('\n'.join(sorted(seen)) + '\n')
        json.dump({'last_rpi': last, 'last_run': dt.date.today().isoformat(), 'last_file': os.path.basename(out)}, open(STATE, 'w'), indent=1)
        print('updated', os.path.basename(SEEN), 'and', os.path.basename(STATE))

if __name__ == '__main__':
    main()
